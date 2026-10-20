import streamlit as st
import pandas as pd
from io import BytesIO
from docx import Document

st.set_page_config(page_title='Аналитика прокуратуры — областной уровень',layout='wide')
st.title('Аналитика прокуратуры РК')
st.caption('Областной уровень • демонстрационный прототип • только вымышленные или разрешённые обезличенные данные')
MODULES=['Надзор за законностью','Уголовный процесс','Защита прав граждан','Представительство интересов государства','Исполнение наказаний и пробация','Координация и иные направления']
METRICS=['Выявленные нарушения','Устранённые нарушения','Обращения граждан','Завершённые производства']
@st.cache_data
def sample():
    rows=[]
    for y in [2025,2026]:
        for ti,t in enumerate(['Город А','Город Б','Район В','Район Г','Район Д']):
            for mi,m in enumerate(MODULES):
                for ki,k in enumerate(METRICS):
                    base=(ti+3)*14+(mi+1)*9+ki*5
                    change=((ti*3+mi+ki)%7-3)*5
                    rows.append({'год':y,'территория':t,'направление':m,'показатель':k,'значение':max(0,base+(change if y==2026 else 0))})
    return pd.DataFrame(rows)
file=st.file_uploader('Загрузить статистику (.xlsx или .csv)',type=['xlsx','csv'])
try:
    df=(pd.read_excel(file) if file.name.lower().endswith('.xlsx') else pd.read_csv(file)) if file else sample()
    required=['год','территория','направление','показатель','значение']
    missing=set(required)-set(df.columns)
    if missing:
        st.error('Отсутствуют столбцы: '+', '.join(sorted(missing)));st.stop()
    df=df[required].copy()
    df['год']=pd.to_numeric(df['год'],errors='coerce')
    df['значение']=pd.to_numeric(df['значение'],errors='coerce')
    bad=df[df[required].isna().any(axis=1)]
    if len(bad): st.warning(f'Строки с пропусками или нечисловыми значениями: {len(bad)}. Они исключены.')
    df=df.dropna(subset=required)
    df['год']=df['год'].astype(int)
    df['территория']=df['территория'].astype(str).str.strip()
    df['направление']=df['направление'].astype(str).str.strip()
    df['показатель']=df['показатель'].astype(str).str.strip()
    dup=df.duplicated(['год','территория','направление','показатель'],keep=False)
    if dup.any():
        st.error(f'Обнаружены дубли показателей для одной территории и года ({dup.sum()} строк). Исправьте файл, чтобы избежать двойного счёта.');st.stop()
    if (df['значение']<0).any(): st.warning('Обнаружены отрицательные значения: проверьте допустимость для выбранного показателя.')
    years=sorted(df['год'].unique())
    year=st.sidebar.selectbox('Отчётный год',years,index=len(years)-1)
    territories=st.sidebar.multiselect('Города и районы',sorted(df['территория'].unique()),default=sorted(df['территория'].unique()))
    modules=st.sidebar.multiselect('Направления',sorted(df['направление'].unique()),default=sorted(df['направление'].unique()))
    metrics=sorted(df['показатель'].unique())
    metric=st.sidebar.selectbox('Сопоставимый показатель',metrics)
    st.sidebar.caption('Для сравнений используются только одинаковые показатели в одинаковых единицах.')
    f=df[(df['год']==year)&df['территория'].isin(territories)&df['направление'].isin(modules)&(df['показатель']==metric)]
    prev=df[(df['год']==year-1)&df['территория'].isin(territories)&df['направление'].isin(modules)&(df['показатель']==metric)]
    if f.empty: st.info('Нет данных для выбранных фильтров.');st.stop()
    curr=f['значение'].sum();old=prev['значение'].sum()
    a,b,c=st.columns(3)
    a.metric(f'{metric}, {year}',f'{curr:,.0f}'.replace(',',' '))
    b.metric('Территорий в выборке',f['территория'].nunique())
    c.metric('Изменение к прошлому году',f'{(curr/old-1)*100:+.1f}%' if old else 'Нет базы сравнения')
    st.subheader('Сравнение территорий по выбранному показателю')
    by_t=f.groupby('территория')['значение'].sum().sort_values(ascending=False)
    st.bar_chart(by_t)
    st.dataframe(by_t.rename('значение').reset_index(),hide_index=True,use_container_width=True)
    st.subheader('Динамика по годам')
    trend=df[df['территория'].isin(territories)&df['направление'].isin(modules)&(df['показатель']==metric)].groupby('год')['значение'].sum()
    st.line_chart(trend)
    st.subheader('Отклонения, требующие проверки')
    comparison=f.groupby('территория')['значение'].sum().to_frame('текущий год')
    comparison['прошлый год']=prev.groupby('территория')['значение'].sum()
    comparison['изменение, %']=((comparison['текущий год']/comparison['прошлый год']-1)*100).round(1)
    comparison.loc[comparison['прошлый год'].isna() | (comparison['прошлый год']==0),'изменение, %']=float('nan')
    comparison['сигнал']=comparison['изменение, %'].apply(lambda x:'Проверить изменение ≥25%' if pd.notna(x) and abs(x)>=25 else ('Нет сопоставимой базы' if pd.isna(x) else 'Без порогового сигнала'))
    st.dataframe(comparison.reset_index(),hide_index=True,use_container_width=True)
    st.caption('Порог 25% — демонстрационный, не нормативный. Сигнал не является выводом о нарушении или оценкой работы прокуратуры.')
    st.subheader('Экспорт')
    st.download_button('Скачать сравнительную таблицу CSV',comparison.reset_index().to_csv(index=False).encode('utf-8-sig'),'regional_comparison.csv','text/csv')
    doc=Document();doc.add_heading('Аналитическая справка (проект)',0)
    doc.add_paragraph(f'Период: {year} год. Показатель: {metric}. Направления: {", ".join(modules)}.')
    doc.add_paragraph(f'Всего по выбранным территориям: {curr:,.0f}. Сравнение с предыдущим годом: '+(f'{(curr/old-1)*100:+.1f}%.' if old else 'нет сопоставимой базы.'))
    doc.add_heading('Сравнение территорий',level=1)
    for t,row in comparison.iterrows():
        delta='нет данных' if pd.isna(row['изменение, %']) else f"{row['изменение, %']:+.1f}%"
        doc.add_paragraph(f'{t}: {row["текущий год"]:,.0f}; изменение: {delta}; {row["сигнал"]}.')
    doc.add_paragraph('Примечание: данные и сигналы требуют проверки сопоставимости, полноты и причин изменений. Автоматические расчёты не заменяют правовую и управленческую оценку.')
    buffer=BytesIO();doc.save(buffer)
    st.download_button('Скачать проект справки Word',buffer.getvalue(),'analytical_report.docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    with st.expander('Формат загрузки и контроль качества'):
        st.write('Обязательные столбцы: год, территория, направление, показатель, значение. Одна строка — один показатель за год, территорию и направление. Не загружайте персональные или секретные сведения.')
        st.dataframe(df.head(10),hide_index=True)
except Exception as exc:
    st.error(f'Ошибка обработки файла: {exc}')
